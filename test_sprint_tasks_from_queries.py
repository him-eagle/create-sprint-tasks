import unittest
from pathlib import Path

from openpyxl import load_workbook

import sprint_tasks_from_queries as sprint_report


REPORT_DIR = Path(__file__).parent / "Reports"
REPORT_TABS = {"All", "FPSO", "Foundation", "CSR"}
TEAM_TABS = REPORT_TABS - {"All"}


class SprintReportTests(unittest.TestCase):
    def test_review_tasks_are_not_design(self):
        """Review tasks must use Review rather than Design / Documentation."""
        self.assertEqual(
            sprint_report.categorize_task("Task", "PR-Review", ""),
            "Review",
        )
        self.assertEqual(
            sprint_report.categorize_task("Task", "Code Review", ""),
            "Review",
        )
        self.assertEqual(
            sprint_report.categorize_task("Task", "Demo Task", ""),
            "Review",
        )

    def test_bugs_keep_bug_fixing_priority(self):
        """Bug work items must remain Bug Fixing despite title keywords."""
        self.assertEqual(
            sprint_report.categorize_task("Bug", "Code Review", ""),
            "Bug Fixing",
        )

    def test_report_scope_accepts_only_query_ids(self):
        """Rows whose PBI and task IDs came from the query must be accepted."""
        rows = [{"pbi_id": 100, "task_id": 101}, {"pbi_id": 100, "task_id": None}]
        sprint_report.validate_report_scope(rows, [100, 101], "FPSO")

    def test_report_scope_rejects_ids_outside_query(self):
        """Rows containing an ID outside the supplied query must be rejected."""
        rows = [{"pbi_id": 100, "task_id": 999}]
        with self.assertRaisesRegex(ValueError, "999"):
            sprint_report.validate_report_scope(rows, [100], "FPSO")

    def test_latest_report_items_are_returned_by_supplied_queries(self):
        """Every item in each latest-report tab must come from its source query."""
        reports = sorted(
            (
                path
                for path in REPORT_DIR.glob("Sprint_Report_*.xlsx")
                if not path.name.startswith("~$")
            ),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if not reports:
            self.skipTest("No generated report exists in Reports")

        report_path = reports[0]
        workbook = load_workbook(report_path, data_only=True, read_only=True)
        pat_token = sprint_report.load_pat_token()
        session = sprint_report.get_session(pat_token) if pat_token else None
        failures = []
        pending_query_tabs = []

        for worksheet in workbook.worksheets:
            if worksheet.title == "All":
                continue
            if worksheet.title not in TEAM_TABS:
                continue

            rows = list(worksheet.iter_rows(values_only=True))
            source_url = rows[0][1] if len(rows[0]) > 1 else None
            if not source_url or not str(source_url).startswith("http"):
                failures.append(
                    {
                        "tab": worksheet.title,
                        "pbi_id": "",
                        "item_id": "",
                        "source_query": source_url or "",
                        "problem": "Source URL is missing or invalid",
                    }
                )
                continue

            try:
                query_info = sprint_report.parse_query_url(source_url)
            except ValueError as error:
                failures.append(
                    {
                        "tab": worksheet.title,
                        "pbi_id": "",
                        "item_id": "",
                        "source_query": source_url,
                        "problem": (
                            "Source is not a saved query URL; expected an Azure DevOps "
                            f"query URL: {error}"
                        ),
                    }
                )
                continue

            if session is None:
                pending_query_tabs.append(worksheet.title)
                continue

            try:
                query_ids = sprint_report.run_wiql_query_by_id(
                    session,
                    query_info["org"],
                    query_info["project"],
                    query_info["query_id"],
                )
            except Exception as error:
                failures.append(
                    {
                        "tab": worksheet.title,
                        "pbi_id": "",
                        "item_id": "",
                        "source_query": source_url,
                        "problem": f"Query error: {error}",
                    }
                )
                continue

            header_index = next(
                index
                for index, row in enumerate(rows)
                if "PBI ID" in row and "Task ID" in row
            )
            headers = rows[header_index]
            pbi_index = headers.index("PBI ID")
            task_index = headers.index("Task ID")
            query_id_set = set(query_ids)
            for row in rows[header_index + 1:]:
                pbi_id = row[pbi_index]
                task_id = row[task_index]
                for item_id in (pbi_id, task_id):
                    if item_id is not None and item_id not in query_id_set:
                        failures.append(
                            {
                                "tab": worksheet.title,
                                "pbi_id": pbi_id or "",
                                "item_id": item_id,
                                "source_query": source_url,
                                "problem": "Item ID was not returned by the supplied query",
                            }
                        )

        if not failures and pending_query_tabs:
            self.skipTest(
                "Set AZURE_DEVOPS_PAT to execute supplied queries for tabs: "
                + ", ".join(pending_query_tabs)
            )

        if failures:
            headers = ["Sheet tab", "PBI ID", "Task/Bug ID", "Source query", "Problem"]
            table_rows = [
                [
                    failure["tab"],
                    failure["pbi_id"],
                    failure["item_id"],
                    failure["source_query"],
                    failure["problem"],
                ]
                for failure in failures
            ]
            widths = [
                max(len(str(value)) for value in [header] + [row[index] for row in table_rows])
                for index, header in enumerate(headers)
            ]
            failure_table = "\n".join(
                " | ".join(str(value).ljust(width) for value, width in zip(row, widths))
                for row in [headers] + table_rows
            )
        else:
            failure_table = ""

        self.assertFalse(
            failures,
            f"Report file '{report_path}' contains items outside the supplied queries.\n"
            f"\n{failure_table}",
        )

    def test_latest_report_assigns_each_pbi_to_one_team(self):
        """The newest workbook must not assign one PBI to multiple team tabs."""
        reports = sorted(
            (
                path
                for path in REPORT_DIR.glob("Sprint_Report_*.xlsx")
                if not path.name.startswith("~$")
            ),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if not reports:
            self.skipTest("No generated report exists in Reports")

        report_path = reports[0]
        workbook = load_workbook(report_path, data_only=True, read_only=True)
        pbi_teams = {}

        for worksheet in workbook.worksheets:
            if worksheet.title not in TEAM_TABS:
                continue

            rows = list(worksheet.iter_rows(values_only=True))
            header_index = next(
                index
                for index, row in enumerate(rows)
                if "Team" in row and "PBI ID" in row
            )
            headers = rows[header_index]
            team_index = headers.index("Team")
            pbi_index = headers.index("PBI ID")

            for row in rows[header_index + 1:]:
                pbi_id = row[pbi_index]
                if pbi_id is None:
                    continue
                self.assertEqual(
                    row[team_index],
                    worksheet.title,
                    f"PBI {pbi_id} has the wrong Team value in {worksheet.title}",
                )
                pbi_teams.setdefault(pbi_id, set()).add(worksheet.title)

        duplicate_assignments = {
            pbi_id: sorted(teams)
            for pbi_id, teams in pbi_teams.items()
            if len(teams) > 1
        }
        self.assertFalse(
            duplicate_assignments,
            f"Report file '{report_path}' assigns PBIs to multiple team tabs: "
            f"{duplicate_assignments}",
        )

    def test_latest_report_categorizes_bugs_as_bug_fixing(self):
        """Every bug in the newest workbook must be categorized as Bug Fixing."""
        reports = sorted(
            (
                path
                for path in REPORT_DIR.glob("Sprint_Report_*.xlsx")
                if not path.name.startswith("~$")
            ),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if not reports:
            self.skipTest("No generated report exists in Reports")

        report_path = reports[0]
        workbook = load_workbook(report_path, data_only=True, read_only=True)
        invalid_bugs = []

        for worksheet in workbook.worksheets:
            if worksheet.title not in TEAM_TABS:
                continue

            rows = list(worksheet.iter_rows(values_only=True))
            header_index = next(
                index
                for index, row in enumerate(rows)
                if "PBI ID" in row and "Task ID" in row
            )
            headers = rows[header_index]
            column_index = {name: headers.index(name) for name in headers if name}

            for row in rows[header_index + 1:]:
                if row[column_index["Task Type"]] != "Bug":
                    continue
                if row[column_index["Task Categorization"]] == "Bug Fixing":
                    continue
                invalid_bugs.append(
                    {
                        "sheet": worksheet.title,
                        "pbi_id": row[column_index["PBI ID"]],
                        "bug_id": row[column_index["Task ID"]],
                        "bug_title": row[column_index["Task Title"]],
                        "category": row[column_index["Task Categorization"]],
                    }
                )

        failure_details = "\n".join(
            (
                f"  Sheet tab: {bug['sheet']} | "
                f"PBI ID: {bug['pbi_id']} | "
                f"Bug ID: {bug['bug_id']} | "
                f"Bug title: {bug['bug_title']} | "
                f"Assigned category: {bug['category']}"
            )
            for bug in invalid_bugs
        )
        self.assertFalse(
            invalid_bugs,
            f"Report file '{report_path}' has bugs outside Bug Fixing.\n"
            f"Each failure is listed below:\n{failure_details}",
        )

    def test_latest_report_categorizes_demo_tasks_as_review(self):
        """Demo tasks in the newest workbook must be categorized as Review."""
        reports = sorted(
            (
                path
                for path in REPORT_DIR.glob("Sprint_Report_*.xlsx")
                if not path.name.startswith("~$")
            ),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if not reports:
            self.skipTest("No generated report exists in Reports")

        report_path = reports[0]
        workbook = load_workbook(report_path, data_only=True, read_only=True)
        demo_tasks = []

        for worksheet in workbook.worksheets:
            if worksheet.title not in TEAM_TABS:
                continue

            rows = list(worksheet.iter_rows(values_only=True))
            header_index = next(
                index
                for index, row in enumerate(rows)
                if "PBI ID" in row and "Task ID" in row
            )
            headers = rows[header_index]
            column_index = {name: headers.index(name) for name in headers if name}

            for row in rows[header_index + 1:]:
                if str(row[column_index["Task Title"]]).strip().lower() != "demo task":
                    continue
                demo_tasks.append(
                    {
                        "sheet": worksheet.title,
                        "pbi_id": row[column_index["PBI ID"]],
                        "task_id": row[column_index["Task ID"]],
                        "task_title": row[column_index["Task Title"]],
                        "category": row[column_index["Task Categorization"]],
                    }
                )

        invalid_demo_tasks = [task for task in demo_tasks if task["category"] != "Review"]
        failure_details = "\n".join(
            (
                f"  Sheet tab: {task['sheet']} | "
                f"PBI ID: {task['pbi_id']} | "
                f"Task ID: {task['task_id']} | "
                f"Task title: {task['task_title']} | "
                f"Assigned category: {task['category']}"
            )
            for task in invalid_demo_tasks
        )
        self.assertFalse(
            invalid_demo_tasks,
            f"Report file '{report_path}' does not categorize demo tasks as Review.\n"
            f"Each failure is listed below:\n{failure_details}",
        )


if __name__ == "__main__":
    unittest.main()
